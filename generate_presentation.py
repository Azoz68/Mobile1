# -*- coding: utf-8 -*-
"""
مولّد عرض تقديمي: الفجوات البحثية — Research Gaps
نسخة موسّعة بشرح تفصيلي وتصميم عصري، مبنية على محتوى ورشة عمل
الدكتورة Hana Hossen Elahemer.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------------------------------------------------------------- palette
NAVY      = RGBColor(0x10, 0x1E, 0x3C)   # deep navy background
NAVY_SOFT = RGBColor(0x1B, 0x2D, 0x55)
INK       = RGBColor(0x18, 0x26, 0x3B)   # main text
MUTED     = RGBColor(0x62, 0x74, 0x8A)   # secondary text
BG        = RGBColor(0xF7, 0xF9, 0xFC)   # light slide background
CARD      = RGBColor(0xFF, 0xFF, 0xFF)
BORDER    = RGBColor(0xE2, 0xE8, 0xF0)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)

TEAL   = RGBColor(0x0E, 0x9F, 0x8C)
GOLD   = RGBColor(0xDD, 0x9F, 0x33)
CORAL  = RGBColor(0xE4, 0x5A, 0x44)
PURPLE = RGBColor(0x6C, 0x63, 0xD8)
PINK   = RGBColor(0xC2, 0x3E, 0x77)
BLUE   = RGBColor(0x2F, 0x6F, 0xB5)
GREEN  = RGBColor(0x3E, 0x8E, 0x5A)

ACCENTS6 = [PINK, GOLD, TEAL, PURPLE, CORAL, GREEN]

FONT = "Dubai"   # ships with Windows/Office; falls back gracefully elsewhere

EMU_W, EMU_H = Inches(13.333), Inches(7.5)
AR_DIGITS = "٠١٢٣٤٥٦٧٨٩"


def tint(color: RGBColor, factor: float) -> RGBColor:
    """Mix color with white. factor=0 -> color, factor=1 -> white."""
    return RGBColor(
        int(color[0] + (255 - color[0]) * factor),
        int(color[1] + (255 - color[1]) * factor),
        int(color[2] + (255 - color[2]) * factor),
    )


def ar_num(n: int) -> str:
    return "".join(AR_DIGITS[int(d)] for d in str(n))


# ---------------------------------------------------------------- helpers
def set_rtl(paragraph):
    paragraph._p.get_or_add_pPr().set("rtl", "1")


def style_run(run, size, color, bold=False, italic=False, font=FONT):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:cs", "a:ea"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def set_alpha(shape, pct):
    """Make a solid-filled shape semi-transparent. pct = opacity 0..100."""
    sp = shape.fill._xPr.find(qn("a:solidFill"))
    clr = sp.find(qn("a:srgbClr"))
    alpha = clr.makeelement(qn("a:alpha"), {"val": str(int(pct * 1000))})
    clr.append(alpha)


def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75,
         shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    no_shadow(s)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    return s


def circle(slide, x, y, d, color, opacity=100):
    s = rect(slide, x, y, d, d, fill=color, shape=MSO_SHAPE.OVAL)
    if opacity < 100:
        set_alpha(s, opacity)
    return s


def textbox(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, wrap=True):
    """paras: list of dicts:
       runs=[(text, size, color, bold, italic)], align, rtl, space_after, line"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, pa in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = pa.get("align", PP_ALIGN.RIGHT)
        if pa.get("rtl", True):
            set_rtl(p)
        p.space_after = Pt(pa.get("space_after", 4))
        p.space_before = Pt(pa.get("space_before", 0))
        p.line_spacing = pa.get("line", 1.12)
        for (t, size, color, *rest) in pa["runs"]:
            bold = rest[0] if len(rest) > 0 else False
            italic = rest[1] if len(rest) > 1 else False
            r = p.add_run()
            r.text = t
            style_run(r, size, color, bold=bold, italic=italic)
    return tb


def P(runs, **kw):
    kw["runs"] = runs
    return kw


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- deck
prs = Presentation()
prs.slide_width, prs.slide_height = EMU_W, EMU_H
BLANK = prs.slide_layouts[6]

MARGIN = 0.55
CW = 13.333 - 2 * MARGIN          # content width
PAGE = [0]                         # page counter


def new_slide(dark=False):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, 13.333, 7.5, fill=(NAVY if dark else BG))
    PAGE[0] += 1
    return s


def footer(slide, accent):
    textbox(slide, MARGIN, 7.06, 3.5, 0.3,
            [P([(f"{PAGE[0]:02d}", 10.5, accent, True),
                ("   Research Gaps — Dr. Hana Hossen Elahemer", 9, MUTED)],
               align=PP_ALIGN.LEFT, rtl=False)])
    textbox(slide, 13.333 - MARGIN - 3.5, 7.06, 3.5, 0.3,
            [P([("الفجوات البحثية • ورشة عمل أكاديمية", 9, MUTED)],
               align=PP_ALIGN.RIGHT)])


def header(slide, icon, title_ar, title_en, accent):
    # left vertical accent strip
    rect(slide, 0, 0, 0.12, 7.5, fill=accent)
    # icon chip (top-right)
    chip = rect(slide, 13.333 - MARGIN - 0.9, 0.45, 0.9, 0.9,
                fill=tint(accent, 0.86), line=accent, line_w=1.2,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.28)
    tf = chip.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = icon
    style_run(r, 30, accent)
    # titles
    tw = CW - 1.1
    textbox(slide, MARGIN, 0.40, tw, 0.62,
            [P([(title_ar, 27, INK, True)], align=PP_ALIGN.RIGHT)])
    textbox(slide, MARGIN, 1.06, tw, 0.4,
            [P([(title_en, 12.5, accent, False, True)],
               align=PP_ALIGN.RIGHT, rtl=False)])
    # underline
    rect(slide, 13.333 - MARGIN - 1.1 - 1.7, 1.52, 1.7, 0.05, fill=accent)
    footer(slide, accent)


# ================================================================ 1 — title
s = new_slide(dark=True)
circle(s, -1.6, -2.1, 4.6, TEAL, 16)
circle(s, 11.2, 4.6, 4.8, PURPLE, 15)
circle(s, 10.4, -1.4, 2.8, GOLD, 18)
circle(s, 1.1, 5.6, 1.5, CORAL, 20)
circle(s, 3.2, 1.05, 0.34, GOLD, 90)
circle(s, 9.85, 1.25, 0.2, TEAL, 90)

textbox(s, 1.5, 1.75, 10.333, 1.35,
        [P([("الفجوات البحثية", 56, GOLD, True)], align=PP_ALIGN.CENTER)])
textbox(s, 1.5, 3.05, 10.333, 0.6,
        [P([("R E S E A R C H   G A P S", 19, WHITE)],
           align=PP_ALIGN.CENTER, rtl=False)])
rect(s, 5.867, 3.78, 1.6, 0.045, fill=GOLD)
textbox(s, 1.5, 4.0, 10.333, 0.55,
        [P([("دليلك الشامل لاكتشاف الفجوات البحثية بذكاء واحترافية",
             17, tint(WHITE, 0), False)], align=PP_ALIGN.CENTER)])
badge = rect(s, 4.79, 4.75, 3.75, 0.52, fill=NAVY_SOFT, line=TEAL, line_w=1,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
tfb = badge.text_frame; tfb.vertical_anchor = MSO_ANCHOR.MIDDLE
pb = tfb.paragraphs[0]; pb.alignment = PP_ALIGN.CENTER; set_rtl(pb)
rb = pb.add_run(); rb.text = "نسخة موسّعة بشرح تفصيلي وتصميم عصري"
style_run(rb, 12.5, TEAL, True)
textbox(s, 1.5, 5.75, 10.333, 0.45,
        [P([("المحتوى العلمي: ", 13, tint(WHITE, 0.0)),
            ("Dr. Hana Hossen Elahemer", 13, GOLD, True)],
           align=PP_ALIGN.CENTER)])
textbox(s, 1.5, 6.25, 10.333, 0.4,
        [P([("ورشة عمل أكاديمية متخصصة • Academic Workshop", 11, MUTED)],
           align=PP_ALIGN.CENTER)])
add_notes(s, "شريحة العنوان: هذا العرض نسخة موسّعة ومعاد تصميمها من ورشة عمل "
             "«الفجوات البحثية» للدكتورة هناء حسين الأهيمر. الهدف: تزويد الباحث "
             "بدليل عملي متكامل لاكتشاف الفجوات البحثية وصياغتها باحترافية، "
             "بدءاً من المفهوم ومروراً بالأنواع والأدوات والاستراتيجيات، "
             "وانتهاءً بخارطة طريق عملية وقالب صياغة جاهز.")

# ================================================================ 2 — agenda
s = new_slide()
header(s, "🗂", "محاور العرض", "Agenda — What you will learn", GOLD)
agenda = [
    ("مفهوم الفجوة البحثية", "التعريف الدقيق والصور الأربع التي تظهر بها الفجوات", PINK),
    ("لماذا تهمّنا الفجوات؟", "ستة أسباب تجعلها حجر الأساس في أي بحث ناجح", GOLD),
    ("الأنواع الستة للفجوات", "من فجوة المعرفة إلى فجوة الزمن — مع سؤال مفتاحي لكل نوع", TEAL),
    ("أمثلة تطبيقية واقعية", "مثال عملي ملموس على كل نوع من الأنواع الستة", PURPLE),
    ("أدوات الذكاء الاصطناعي", "ثماني أدوات ذكية تسرّع اكتشاف الفجوات وتوثيقها", CORAL),
    ("استراتيجيات المحترفين", "أربع استراتيجيات ميدانية + التفكير النقدي عند التحليل", GREEN),
    ("من الفجوة إلى الصياغة", "تمييز الفجوة القوية، خارطة طريق عملية، وقالب صياغة جاهز", BLUE),
]
cw, gap = (CW - 0.5) / 2, 0.5
rows_y = 1.85
for i, (t, d, ac) in enumerate(agenda):
    col = i // 4          # 0 = right column (items 1-4), 1 = left
    row = i % 4
    x = MARGIN + (cw + gap) * (1 - col) if False else (13.333 - MARGIN - cw if col == 0 else MARGIN)
    y = rows_y + row * 1.26
    card = rect(s, x, y, cw, 1.1, fill=CARD, line=BORDER, line_w=1,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
    # number chip on the right
    nc = rect(s, x + cw - 0.75, y + 0.24, 0.62, 0.62, fill=tint(ac, 0.85),
              line=ac, line_w=1, shape=MSO_SHAPE.OVAL)
    tfn = nc.text_frame; tfn.vertical_anchor = MSO_ANCHOR.MIDDLE
    pn = tfn.paragraphs[0]; pn.alignment = PP_ALIGN.CENTER
    rn = pn.add_run(); rn.text = ar_num(i + 1); style_run(rn, 16, ac, True)
    textbox(s, x + 0.25, y + 0.13, cw - 1.15, 0.45,
            [P([(t, 14.5, INK, True)], align=PP_ALIGN.RIGHT)])
    textbox(s, x + 0.25, y + 0.58, cw - 1.15, 0.45,
            [P([(d, 10.5, MUTED)], align=PP_ALIGN.RIGHT)])
add_notes(s, "محاور العرض السبعة: نبدأ بالمفهوم ثم الأهمية، وننتقل إلى الأنواع "
             "الستة مع أمثلة تطبيقية، ثم أدوات الذكاء الاصطناعي المساعدة، "
             "فاستراتيجيات الباحثين المحترفين والتفكير النقدي، ونختم بكيفية "
             "التمييز بين الفجوة القوية والضعيفة وخطوات عملية للوصول إلى "
             "صياغة فجوة بحثية جاهزة للاستخدام في المقترح البحثي.")

# ================================================================ 3 — definition
s = new_slide()
header(s, "🧠", "ما هي الفجوة البحثية؟", "What is a Research Gap?", PINK)
# definition panel
panel = rect(s, MARGIN, 1.8, CW, 1.72, fill=tint(PURPLE, 0.92), line=PURPLE,
             line_w=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
textbox(s, MARGIN + 0.35, 1.95, CW - 0.7, 1.5,
        [P([("الفجوة البحثية (Research Gap): ", 15, PURPLE, True),
            ("هي المنطقة أو المشكلة في مجال معرفي لم تُدرس بعد، أو دُرست بصورة "
             "غير كافية، أو تحتاج إلى مزيد من الاستكشاف والتعمق من قِبل الباحثين.",
             13.5, INK)], align=PP_ALIGN.RIGHT, space_after=7, line=1.25),
         P([("بعبارة أبسط: ", 12.5, PINK, True),
            ("هي «المساحة الفارغة» بين ما تعرفه البشرية اليوم وما تحتاج إلى "
             "معرفته غداً — وسدّ هذه المساحة هو جوهر أي بحث علمي أصيل.",
             12.5, INK)], align=PP_ALIGN.RIGHT, line=1.25)])
# four forms
forms = [
    ("الغياب الكلّي", "🔍", "موضوع لم يتناوله أي باحث من قبل — أندر الصور وأثمنها متى كان الموضوع مهماً", PINK),
    ("الدراسة الجزئية", "⚡", "موضوع دُرس من زاوية واحدة فقط بينما أُهملت زواياه الأخرى", GOLD),
    ("تناقض النتائج", "🔄", "دراسات وصلت إلى نتائج متضاربة تستدعي التحقيق والحسم", TEAL),
    ("السياق الجغرافي", "🌍", "نتائج لم تُختبر في بيئتك أو مجتمعك أو ثقافتك المحلية", GREEN),
]
cw4 = (CW - 3 * 0.28) / 4
for i, (t, ic, d, ac) in enumerate(forms):
    x = 13.333 - MARGIN - cw4 - i * (cw4 + 0.28)
    y = 3.75
    rect(s, x, y, cw4, 3.0, fill=CARD, line=BORDER, line_w=1,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    rect(s, x, y, cw4, 0.14, fill=ac, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=1.0)
    textbox(s, x + 0.15, y + 0.35, cw4 - 0.3, 0.6,
            [P([(ic, 24, ac)], align=PP_ALIGN.CENTER, rtl=False)])
    textbox(s, x + 0.15, y + 1.0, cw4 - 0.3, 0.5,
            [P([(t, 15, ac, True)], align=PP_ALIGN.CENTER)])
    textbox(s, x + 0.2, y + 1.55, cw4 - 0.4, 1.3,
            [P([(d, 11, INK)], align=PP_ALIGN.CENTER, line=1.22)])
add_notes(s, "التعريف: الفجوة البحثية هي منطقة معرفية لم تُدرس أو دُرست بشكل "
             "ناقص. تظهر بأربع صور رئيسية: (1) الغياب الكلي: لم يتناول الموضوعَ "
             "أحد إطلاقاً، وهي أندر الحالات وأعلاها قيمة. (2) الدراسة الجزئية: "
             "الموضوع مدروس من زاوية واحدة فقط (مثلاً: أثر تقني دون الأثر "
             "النفسي). (3) تناقض النتائج: دراسات متضاربة تحتاج دراسة حاسمة. "
             "(4) السياق الجغرافي/الثقافي: نتائج مثبتة في مجتمعات أخرى لكنها "
             "غير مختبرة في مجتمعك — وهذه أكثر الفجوات شيوعاً وواقعية للباحث العربي.")

# ================================================================ 4 — why matter
s = new_slide()
header(s, "⚙️", "لماذا تهمّنا الفجوات البحثية؟", "Why Research Gaps Matter", GOLD)
why = [
    ("توجيه البحث", "🧭", "تمنحك اتجاهاً واضحاً لرسالتك أو أطروحتك بدلاً من التخبط بين مواضيع مكررة ومستهلكة", PINK),
    ("الأصالة العلمية", "🏆", "تضمن أن بحثك يُضيف جديداً حقيقياً للمعرفة ولا يعيد إنتاج ما هو موجود أصلاً", GOLD),
    ("تبرير الدراسة", "⚖️", "حجّتك الأقوى لإقناع اللجنة العلمية والمشرفين بقيمة بحثك وجدوى إجرائه", PURPLE),
    ("ربط الأدبيات", "🔗", "تبني جسراً منطقياً بين ما هو موجود في الأدبيات وما تطمح لاكتشافه بنفسك", TEAL),
    ("النشر الأكاديمي", "📰", "المجلات المحكّمة تفضّل الأبحاث التي تعالج فجوات واضحة ومعلَنة منذ المقدمة", GREEN),
    ("التمويل البحثي", "💰", "الجهات الممولة تبحث عن مقترحات تسدّ فجوات حقيقية ذات أثر ملموس", CORAL),
]
cw3, gap3 = (CW - 2 * 0.28) / 3, 0.28
for i, (t, ic, d, ac) in enumerate(why):
    col, row = i % 3, i // 3
    x = 13.333 - MARGIN - cw3 - col * (cw3 + gap3)
    y = 1.85 + row * 2.6
    rect(s, x, y, cw3, 2.38, fill=CARD, line=BORDER, line_w=1,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    rect(s, x + cw3 - 0.09, y + 0.25, 0.09, 1.9, fill=ac)
    textbox(s, x + 0.25, y + 0.22, cw3 - 0.55, 0.5,
            [P([(t + "  ", 14.5, ac, True), (ic, 13, ac)], align=PP_ALIGN.RIGHT)])
    textbox(s, x + 0.25, y + 0.78, cw3 - 0.55, 1.45,
            [P([(d, 11.5, INK)], align=PP_ALIGN.RIGHT, line=1.28)])
add_notes(s, "ستة أسباب تجعل تحديد الفجوة أهم خطوة في البحث: (1) توجيه البحث: "
             "بوصلة تحدد وجهة الرسالة وتوفر شهوراً من التخبط. (2) الأصالة: شرط "
             "القبول في الدراسات العليا والمجلات — بحث بلا فجوة = تكرار. "
             "(3) التبرير: أول سؤال تطرحه لجنة المناقشة: «لماذا هذا البحث؟» "
             "والفجوة الموثقة هي الجواب. (4) ربط الأدبيات: تُظهر تمكنك من "
             "المجال وقراءتك النقدية. (5) النشر: المحكّمون يبحثون عن عبارة "
             "«تسد هذه الدراسة فجوة…» في المقدمة. (6) التمويل: الممولون "
             "يستثمرون في الجديد المؤثر لا في المكرر.")

# ================================================================ 5 — six types
s = new_slide()
header(s, "🔎", "الأنواع الستة للفجوات البحثية", "Six Types of Research Gaps", TEAL)
types = [
    ("فجوة المعرفة", "Knowledge Gap", "💡",
     "معلومة أو ظاهرة لا يعرفها العلم بعد في مجال معيّن",
     "السؤال: ماذا لا نعرف بعد؟", PINK),
    ("فجوة الممارسة", "Practice Gap", "⚙️",
     "الفجوة بين ما تقوله الأبحاث وما يحدث فعلاً في الميدان",
     "السؤال: لماذا لا يُطبَّق ما نعرفه؟", GOLD),
    ("فجوة الأدلة", "Evidence Gap", "🧪",
     "نقص الأدلة التجريبية اللازمة لدعم فرضية أو ادعاء معيّن",
     "السؤال: أين الدليل العلمي؟", TEAL),
    ("فجوة السكان", "Population Gap", "👥",
     "مجموعة لم تُدرس بكفاية: جنس، عمر، ثقافة، أو فئة مهنية",
     "السؤال: مَن الذي أُهمل من الدراسة؟", PURPLE),
    ("فجوة المنهجية", "Methodological Gap", "🧩",
     "مناهج مختلفة قد تكشف نتائج جديدة تماماً عن نفس الموضوع",
     "السؤال: ماذا لو بحثنا بطريقة أخرى؟", CORAL),
    ("فجوة الزمن", "Temporal Gap", "⏳",
     "دراسات قديمة لم تعد تواكب التغيرات والمستجدات الحديثة",
     "السؤال: هل ما زالت النتائج صحيحة اليوم؟", GREEN),
]
for i, (t, en, ic, d, q, ac) in enumerate(types):
    col, row = i % 3, i // 3
    x = 13.333 - MARGIN - cw3 - col * (cw3 + gap3)
    y = 1.85 + row * 2.6
    rect(s, x, y, cw3, 2.38, fill=tint(ac, 0.93), line=ac, line_w=1.1,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    nc = rect(s, x + cw3 - 0.72, y + 0.2, 0.52, 0.52, fill=ac, shape=MSO_SHAPE.OVAL)
    tfn = nc.text_frame; tfn.vertical_anchor = MSO_ANCHOR.MIDDLE
    pn = tfn.paragraphs[0]; pn.alignment = PP_ALIGN.CENTER
    rn = pn.add_run(); rn.text = ar_num(i + 1); style_run(rn, 14, WHITE, True)
    textbox(s, x + 0.22, y + 0.16, cw3 - 1.0, 0.42,
            [P([(t, 14, ac, True)], align=PP_ALIGN.RIGHT)])
    textbox(s, x + 0.22, y + 0.58, cw3 - 1.0, 0.32,
            [P([(en + "  " + ic, 10.5, MUTED, False, True)],
               align=PP_ALIGN.RIGHT, rtl=False)])
    textbox(s, x + 0.24, y + 0.98, cw3 - 0.48, 0.85,
            [P([(d, 11, INK)], align=PP_ALIGN.RIGHT, line=1.22)])
    textbox(s, x + 0.24, y + 1.88, cw3 - 0.48, 0.42,
            [P([(q, 10.5, ac, True)], align=PP_ALIGN.RIGHT)])
add_notes(s, "الأنواع الستة: (1) فجوة المعرفة: غياب معلومة كلياً — الأكثر "
             "كلاسيكية. (2) فجوة الممارسة/التطبيق: العلم يقول شيئاً والواقع "
             "يفعل غيره؛ مثالية للتخصصات المهنية كالتمريض والتعليم والإدارة. "
             "(3) فجوة الأدلة: توصيات شائعة دون تجارب مضبوطة تدعمها. "
             "(4) فجوة السكان: فئات مهمَلة (نساء، كبار سن، مجتمعات عربية…). "
             "(5) فجوة المنهجية: الموضوع دُرس كمياً فقط أو نوعياً فقط، وتغيير "
             "المنهج قد يقلب الصورة. (6) فجوة الزمن: نتائج ما قبل تحولات كبرى "
             "(الجائحة، الذكاء الاصطناعي) تحتاج إعادة اختبار. احفظ السؤال "
             "المفتاحي لكل نوع — فهو أسرع طريقة لاستدعائها عند القراءة.")

# ================================================================ 6 — examples
s = new_slide()
header(s, "📚", "أمثلة تطبيقية لكل نوع من الفجوات", "Practical, Real-World Examples", PURPLE)
examples = [
    ("معرفة", "لم تُدرس تأثيرات الذكاء الاصطناعي على القرارات الأخلاقية في القطاع الصحي العربي",
     "غياب كامل لمعرفة مطلوبة في مجال حسّاس وحديث", PINK),
    ("ممارسة", "البروتوكولات الدولية للتمريض لا تُطبَّق كما ينبغي في المستشفيات الإقليمية",
     "المعرفة موجودة لكن التطبيق الميداني غائب أو مشوَّه", GOLD),
    ("أدلة", "توصيات كثيرة بدمج التعلم المصغّر في التدريب دون تجارب مضبوطة تُثبت فاعليته",
     "ادعاء شائع ينقصه الدليل التجريبي الصارم", TEAL),
    ("سكانية", "معظم دراسات القلق الاجتماعي أُجريت على طلاب غربيين ولا تمثّل الشباب العربي",
     "العينات المدروسة لا تعكس مجتمعك المستهدف", PURPLE),
    ("منهجية", "الدراسات استخدمت المنهج الكمي فقط؛ والمنهج النوعي قد يكشف تفاصيل أعمق بكثير",
     "تغيير أسلوب البحث يفتح نافذة جديدة على الظاهرة", CORAL),
    ("زمنية", "أبحاث التجارة الإلكترونية المتاحة تعود لما قبل 2020 ولا تعكس أثر الجائحة",
     "الواقع تغيّر جذرياً والدراسات لم تلحق به بعد", GREEN),
]
for i, (tag, ex, why_g, ac) in enumerate(examples):
    y = 1.82 + i * 0.87
    rect(s, MARGIN, y, CW, 0.75, fill=tint(ac, 0.94), line=tint(ac, 0.45),
         line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.22)
    tg = rect(s, 13.333 - MARGIN - 1.35, y + 0.13, 1.2, 0.49, fill=ac,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tft = tg.text_frame; tft.vertical_anchor = MSO_ANCHOR.MIDDLE
    pt = tft.paragraphs[0]; pt.alignment = PP_ALIGN.CENTER; set_rtl(pt)
    rt = pt.add_run(); rt.text = tag; style_run(rt, 11.5, WHITE, True)
    textbox(s, MARGIN + 0.3, y + 0.075, CW - 1.95, 0.35,
            [P([(ex, 11.5, INK, True)], align=PP_ALIGN.RIGHT)])
    textbox(s, MARGIN + 0.3, y + 0.42, CW - 1.95, 0.3,
            [P([("لماذا هي فجوة؟  ", 9.5, ac, True), (why_g, 9.5, MUTED)],
               align=PP_ALIGN.RIGHT)])
add_notes(s, "أمثلة واقعية توضح كل نوع: لاحظ أن كل مثال يحدد (الموضوع + "
             "السياق + سبب كونه فجوة). مثال المعرفة: الذكاء الاصطناعي "
             "والقرارات الأخلاقية في الصحة العربية — موضوع حديث لم يُبحث. "
             "مثال الممارسة: بروتوكولات التمريض معروفة عالمياً لكنها لا تُطبق "
             "إقليمياً — الفجوة في التطبيق لا المعرفة. مثال الأدلة: التعلم "
             "المصغّر يُوصى به دون تجارب مضبوطة. مثال السكان: عينات غربية "
             "لا تمثل الشباب العربي. مثال المنهجية: الاقتصار على الكمي يحجب "
             "التفاصيل النوعية. مثال الزمن: بيانات ما قبل 2020 لم تعد تعكس "
             "واقع التجارة الإلكترونية. استخدم البنية نفسها عند صياغة فجوتك.")

# ================================================================ 7 — AI tools
s = new_slide()
header(s, "🤖", "أدوات الذكاء الاصطناعي لاكتشاف الفجوات",
       "AI Tools for Discovering Research Gaps", CORAL)
tools = [
    ("Elicit", "🔍", "يلخّص الأبحاث ويستخرج النتائج والجداول بسؤال مباشر بلغة طبيعية", PINK),
    ("Consensus", "🤝", "يرصد إجماع الأبحاث حول سؤال علمي محدد ويُظهر اتجاه الأدلة", TEAL),
    ("Research Rabbit", "🐇", "يرسم خريطة علاقات المقالات والباحثين بصرياً لتتبع تطور المجال", GOLD),
    ("Connected Papers", "🕸", "شبكة بصرية للاقتباسات تكشف الدراسات المرتبطة ببحثك المحوري", PURPLE),
    ("Scite.ai", "📊", "يميّز الأبحاث الداعمة والمعارضة لأي ادعاء علمي (اقتباسات ذكية)", CORAL),
    ("Semantic Scholar", "🧠", "محرك بحث دلالي يفهم المعنى والسياق لا الكلمات المفتاحية فقط", GREEN),
    ("Claude", "⭐", "يحلّل الأدبيات بعمق ويقترح فجوات محتملة عبر حوار نقدي متدرج", NAVY_SOFT),
    ("ChatGPT", "💬", "يساعد في صياغة أسئلة البحث واكتشاف المحاور الغائبة في خطتك", BLUE),
]
cw4b, gap4 = (CW - 3 * 0.24) / 4, 0.24
for i, (name, ic, d, ac) in enumerate(tools):
    col, row = i % 4, i // 4
    x = 13.333 - MARGIN - cw4b - col * (cw4b + gap4)
    y = 1.82 + row * 2.12
    rect(s, x, y, cw4b, 1.95, fill=CARD, line=BORDER, line_w=1,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.11)
    rect(s, x, y, cw4b, 0.12, fill=ac, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=1.0)
    textbox(s, x + 0.15, y + 0.24, cw4b - 0.3, 0.42,
            [P([(ic + "  " + name, 13.5, ac, True)],
               align=PP_ALIGN.CENTER, rtl=False)])
    textbox(s, x + 0.18, y + 0.72, cw4b - 0.36, 1.15,
            [P([(d, 10, INK)], align=PP_ALIGN.CENTER, line=1.2)])
tip = rect(s, MARGIN, 6.15, CW, 0.62, fill=tint(GOLD, 0.88), line=GOLD,
           line_w=1.1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
textbox(s, MARGIN + 0.3, 6.27, CW - 0.6, 0.4,
        [P([("💡 قاعدة ذهبية: ", 12, GOLD, True),
            ("هذه الأدوات نقطة انطلاق تسرّع البحث، وليست حكماً نهائياً — تحقّق دائماً من المصادر الأصلية قبل الاستشهاد بأي نتيجة.",
             11.5, INK)], align=PP_ALIGN.CENTER)])
add_notes(s, "ثماني أدوات ذكاء اصطناعي تخدم مراحل مختلفة: للاستطلاع السريع "
             "وتلخيص الأدلة: Elicit وConsensus. لرسم خريطة المجال بصرياً "
             "واكتشاف العناقيد والفراغات: Research Rabbit وConnected Papers. "
             "لفحص جودة الأدلة واتجاهها: Scite.ai يميز الاقتباسات الداعمة من "
             "المعارضة — ممتاز لاكتشاف التناقضات. للبحث العميق: Semantic "
             "Scholar محرك دلالي أكاديمي. للتحليل والحوار النقدي وصياغة "
             "الفجوات: Claude وChatGPT. سير عمل مقترح: ابدأ بـ Semantic "
             "Scholar لجمع الأوراق، ثم Connected Papers للخريطة، ثم Elicit "
             "لاستخراج النتائج، واختم بحوار نقدي مع Claude لصياغة الفجوة. "
             "وتذكّر: الأداة تقترح والباحث يتحقق ويقرر.")

# ================================================================ 8 — strategies
s = new_slide()
header(s, "🗞", "استراتيجيات الباحثين المحترفين", "Expert Researcher Strategies", GREEN)
strats = [
    ("قراءة خاتمات الدراسات", "📄",
     "الباحثون يوصون بفجوات مستقبلية في نهاية أبحاثهم ضمن قسم «حدود الدراسة والبحوث المقترحة» — استثمرها فهي فجوات جاهزة وموثّقة!",
     PURPLE),
    ("رسم خريطة الأدبيات", "🗺",
     "صنّف الدراسات في جدول: السكان، المنهج، البلد، الزمن — ستظهر الخانات الفارغة جليّاً وكل خانة فارغة فجوة محتملة.",
     TEAL),
    ("البحث في التناقضات", "⚖️",
     "التضارب بين نتائج الدراسات = فجوة قوية وفرصة بحثية ذهبية؛ دراستك قد تكون الحكم الذي يحسم الخلاف العلمي.",
     PINK),
    ("اسأل: ماذا لو؟", "❓",
     "ماذا لو طبّقنا هذه الدراسة على سياق مختلف أو فئة عمرية أخرى أو بمنهج مغاير؟ كل إجابة مقنعة قد تكون بحثاً جديداً.",
     GOLD),
]
for i, (t, ic, d, ac) in enumerate(strats):
    y = 1.82 + i * 1.06
    rect(s, MARGIN, y, CW, 0.93, fill=tint(ac, 0.93), line=ac, line_w=1.1,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16)
    icr = rect(s, 13.333 - MARGIN - 0.85, y + 0.16, 0.62, 0.62, fill=ac,
               shape=MSO_SHAPE.OVAL)
    tfi = icr.text_frame; tfi.vertical_anchor = MSO_ANCHOR.MIDDLE
    pi = tfi.paragraphs[0]; pi.alignment = PP_ALIGN.CENTER
    ri = pi.add_run(); ri.text = ic; style_run(ri, 17, WHITE)
    textbox(s, MARGIN + 0.3, y + 0.08, CW - 1.35, 0.4,
            [P([(t, 13.5, ac, True)], align=PP_ALIGN.RIGHT)])
    textbox(s, MARGIN + 0.3, y + 0.47, CW - 1.35, 0.42,
            [P([(d, 10.8, INK)], align=PP_ALIGN.RIGHT, line=1.15)])
tip2 = rect(s, MARGIN, 6.12, CW, 0.65, fill=NAVY, line=GOLD, line_w=1.2,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
textbox(s, MARGIN + 0.3, 6.25, CW - 0.6, 0.42,
        [P([("💡 نصيحة ذهبية: ", 12.5, GOLD, True),
            ("اقرأ آخر ٢٠ ورقة في مجلتك المستهدفة — الأنماط المتكررة ستكشف الفراغات بوضوح تام.",
             12, WHITE)], align=PP_ALIGN.CENTER)])
add_notes(s, "أربع استراتيجيات ميدانية يعتمدها المحترفون: (1) خاتمات الدراسات: "
             "قسم Limitations & Future Research هو منجم فجوات جاهزة أوصى بها "
             "أصحابها — أسرع طريق للمبتدئين. (2) خريطة الأدبيات: جدول "
             "«مصفوفة الأدبيات» بأعمدة السكان/المنهج/البلد/الزمن يجعل الفراغ "
             "بصرياً لا تخمينياً. (3) التناقضات: عندما تجد دراسة تقول (نعم) "
             "وأخرى تقول (لا)، فهناك متغير وسيط لم يُكتشف — فرصتك. "
             "(4) سؤال «ماذا لو»: أداة توليد الأفكار الأقوى؛ بدّل السياق أو "
             "الفئة أو المنهج واختبر منطقية السؤال الجديد. النصيحة الذهبية: "
             "قراءة آخر 20 ورقة في المجلة المستهدفة تكشف ذوق المجلة "
             "واتجاهاتها والفراغات التي يتحدث عنها الجميع.")

# ================================================================ 9 — critical thinking
s = new_slide()
header(s, "💡", "التفكير النقدي عند تحليل الدراسات", "Critical Thinking Framework", BLUE)
pw = (CW - 0.35) / 2
px_r = 13.333 - MARGIN - pw          # right panel x
px_l = MARGIN                         # left panel x
# right panel — six golden questions
rect(s, px_r, 1.82, pw, 4.15, fill=tint(PURPLE, 0.93), line=PURPLE, line_w=1.1,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
textbox(s, px_r + 0.3, 1.98, pw - 0.6, 0.45,
        [P([("❓ الأسئلة الستة الذهبية", 15, PURPLE, True)], align=PP_ALIGN.RIGHT)])
qs = [
    ("مَن", "دُرس؟ (السكان والعينة)", PINK),
    ("ماذا", "قِيس؟ (المتغيرات والأدوات)", GOLD),
    ("أين", "أُجري؟ (السياق الجغرافي والثقافي)", TEAL),
    ("متى", "نُشر؟ (الزمن ومدى الحداثة)", PURPLE),
    ("كيف", "بُحث؟ (المنهجية والتصميم)", GREEN),
    ("لماذا", "توقّف؟ (حدود الدراسة المعلنة)", CORAL),
]
qparas = []
for w, restq, ac in qs:
    qparas.append(P([("●  ", 11, ac), (w + " ", 12.5, ac, True),
                     (restq, 12, INK)],
                    align=PP_ALIGN.RIGHT, space_after=9, line=1.15))
textbox(s, px_r + 0.35, 2.55, pw - 0.7, 3.3, qparas)
# left panel — researchable gap indicators
rect(s, px_l, 1.82, pw, 4.15, fill=tint(TEAL, 0.93), line=TEAL, line_w=1.1,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
textbox(s, px_l + 0.3, 1.98, pw - 0.6, 0.45,
        [P([("🔎 مؤشرات الفجوة القابلة للبحث", 15, TEAL, True)], align=PP_ALIGN.RIGHT)])
inds = [
    "«لم يُدرَس هذا في …» — غياب صريح يذكره الباحثون",
    "«الأبحاث تفترض ولم تختبر …» — ادعاء بلا دليل",
    "«النتائج متضاربة حول …» — تناقض يستدعي الحسم",
    "«العينة لا تمثّل …» — قصور في التعميم",
    "«المنهج الكمي وحده لا يكفي …» — حاجة منهجية",
    "«الدراسات تجاهلت متغير …» — عامل مفقود من النموذج",
]
iparas = [P([("✔  ", 11.5, TEAL, True), (t, 11.5, INK)],
             align=PP_ALIGN.RIGHT, space_after=9, line=1.15) for t in inds]
textbox(s, px_l + 0.35, 2.55, pw - 0.7, 3.3, iparas)
# quote bar
rect(s, MARGIN, 6.2, CW, 0.58, fill=tint(PINK, 0.9), line=PINK, line_w=1,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
textbox(s, MARGIN + 0.3, 6.31, CW - 0.6, 0.38,
        [P([("«التفكير النقدي هو البناء على أكتاف العمالقة لرؤية ما وراء الأفق»",
             12.5, PINK, True)], align=PP_ALIGN.CENTER)])
add_notes(s, "إطار عمل التفكير النقدي عند قراءة أي دراسة: اطرح الأسئلة الستة "
             "الذهبية — مَن دُرس (هل العينة تشبه مجتمعك؟)، ماذا قيس (هل هناك "
             "متغيرات مُهملة؟)، أين أُجري (هل السياق مختلف عن سياقك؟)، متى "
             "نُشر (هل تجاوزتها الأحداث؟)، كيف بُحث (هل يصلح منهج آخر؟)، "
             "لماذا توقفت الدراسة عند هذا الحد (ماذا تقول حدودها؟). وأثناء "
             "القراءة التقط العبارات المؤشرة على فجوة قابلة للبحث: «لم يُدرس»، "
             "«نتائج متضاربة»، «العينة لا تمثل»، «تجاهلت متغير…». كل عبارة "
             "منها بذرة سؤال بحثي جديد. اجعل هذه الأسئلة قائمة فحص ثابتة "
             "بجوارك عند مراجعة الأدبيات.")

# ================================================================ 10 — strong vs weak
s = new_slide()
header(s, "⚖️", "كيف تميّز الفجوة القوية من الضعيفة؟", "Strong Gap vs. Weak Gap", TEAL)
# right = strong
rect(s, px_r, 1.82, pw, 3.95, fill=tint(TEAL, 0.93), line=TEAL, line_w=1.2,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
textbox(s, px_r + 0.3, 1.98, pw - 0.6, 0.45,
        [P([("💪 الفجوة القوية", 15.5, TEAL, True)], align=PP_ALIGN.RIGHT)])
strong = [
    ("مدعومة بأدلة من الأدبيات", "تستشهد بدراسات حقيقية تثبت الغياب"),
    ("محددة وقابلة للقياس", "موضوع + فئة + سياق + زمن واضح"),
    ("ذات أهمية علمية وعملية", "سدّها يغيّر فهماً أو ممارسة"),
    ("معالجتها ممكنة بمنهجية واضحة", "تصميم بحثي واقعي وقابل للتنفيذ"),
    ("مرتبطة بسياق محدد", "بيئة أو مجتمع أو قطاع معيّن"),
]
sparas = []
for t, d in strong:
    sparas.append(P([("✔  ", 12, TEAL, True), (t, 12, INK, True),
                     ("  —  " + d, 10.5, MUTED)],
                    align=PP_ALIGN.RIGHT, space_after=8, line=1.15))
textbox(s, px_r + 0.35, 2.55, pw - 0.7, 3.1, sparas)
# left = weak
rect(s, px_l, 1.82, pw, 3.95, fill=tint(CORAL, 0.93), line=CORAL, line_w=1.2,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
textbox(s, px_l + 0.3, 1.98, pw - 0.6, 0.45,
        [P([("⚠️ الفجوة الضعيفة", 15.5, CORAL, True)], align=PP_ALIGN.RIGHT)])
weak = [
    ("عامة وفضفاضة جداً", "بلا موضوع محدد أو سياق واضح"),
    ("غير مدعومة بمراجعة كافية", "انطباع شخصي لا مسح منهجي"),
    ("«لم يدرسها أحد» بلا دليل", "ادعاء خطير يسهل دحضه أمام اللجنة"),
    ("لا تُضيف قيمة للمعرفة", "سدّها لا يغيّر شيئاً يُذكر"),
    ("تعيد ما قالته دراسات سابقة", "تكرار مقنَّع وليس بحثاً جديداً"),
]
wparas = []
for t, d in weak:
    wparas.append(P([("✘  ", 12, CORAL, True), (t, 12, INK, True),
                     ("  —  " + d, 10.5, MUTED)],
                    align=PP_ALIGN.RIGHT, space_after=8, line=1.15))
textbox(s, px_l + 0.35, 2.55, pw - 0.7, 3.1, wparas)
# comparative example
rect(s, MARGIN, 5.95, CW, 0.88, fill=tint(GOLD, 0.9), line=GOLD, line_w=1.1,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
textbox(s, MARGIN + 0.3, 6.03, CW - 0.6, 0.75,
        [P([("🧪 مثال مقارن —  ", 11.5, GOLD, True),
            ("ضعيفة: ", 11.5, CORAL, True),
            ("«لم تُدرس الصحة النفسية بشكل كافٍ»", 11.5, INK)],
           align=PP_ALIGN.RIGHT, space_after=3),
         P([("قوية: ", 11.5, TEAL, True),
            ("«لا توجد دراسة تقيس أثر الدراسة عن بُعد على القلق لدى طلاب الجامعات العربية بعد 2020»",
             11.5, INK, True)], align=PP_ALIGN.RIGHT)])
add_notes(s, "معيار الجودة: الفجوة القوية موثقة (تستند إلى مراجعة أدبيات "
             "فعلية)، محددة (موضوع + فئة + سياق + إطار زمني)، مهمة (سدها "
             "يغير فهماً علمياً أو ممارسة عملية)، قابلة للمعالجة بمنهجية "
             "واقعية، ومرتبطة بسياق واضح. الفجوة الضعيفة عكس ذلك: فضفاضة، "
             "انطباعية، تدّعي «لم يدرسها أحد» دون دليل — وهذه العبارة تحديداً "
             "خطيرة لأن أي عضو لجنة يجد دراسة واحدة مضادة يهدم بها مصداقيتك. "
             "لاحظ الفرق في المثال المقارن: العبارة الضعيفة عامة بلا حدود، "
             "بينما القوية تحدد المتغير (القلق) والفئة (طلاب الجامعات العربية) "
             "والسياق (الدراسة عن بعد) والزمن (بعد 2020) — جاهزة لتتحول إلى "
             "سؤال بحثي مباشرة.")

# ================================================================ 11 — roadmap
s = new_slide()
header(s, "🚀", "خارطة طريق عملية: من الصفر إلى الفجوة", "A 7-Step Practical Roadmap", CORAL)
steps = [
    ("حدّد مجالك بدقة", "اختر تخصصاً فرعياً ضيقاً — «تعليم إلكتروني للمرحلة الابتدائية» لا «التعليم» عموماً", PINK),
    ("اجمع أحدث الدراسات", "20–30 دراسة من آخر خمس سنوات عبر Semantic Scholar وGoogle Scholar", GOLD),
    ("اقرأ الخاتمات أولاً", "ابدأ بأقسام «حدود الدراسة» و«البحوث المقترحة» — أسرع مصدر للفجوات الموثقة", TEAL),
    ("أنشئ مصفوفة الأدبيات", "جدول يقارن: العينة، السياق، المنهج، الزمن، أبرز النتائج لكل دراسة", PURPLE),
    ("علّم الفراغات والتناقضات", "ظلّل الخانات الفارغة والنتائج المتضاربة — هذه مواضع الفجوات المحتملة", CORAL),
    ("تحقّق بأدوات الذكاء الاصطناعي", "استخدم Consensus وConnected Papers للتأكد أن الفجوة لم تُسدّ حديثاً", GREEN),
    ("صُغ فجوتك بدقة", "جملة محددة قابلة للقياس تذكر الموضوع والفئة والسياق والزمن", BLUE),
]
col_w = (CW - 0.4) / 2
for i, (t, d, ac) in enumerate(steps):
    col = 0 if i < 4 else 1              # 0 = right column
    row = i if i < 4 else i - 4
    x = 13.333 - MARGIN - col_w if col == 0 else MARGIN
    y = 1.85 + row * 1.24
    rect(s, x, y, col_w, 1.08, fill=CARD, line=BORDER, line_w=1,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16)
    nc = rect(s, x + col_w - 0.82, y + 0.21, 0.64, 0.64, fill=ac,
              shape=MSO_SHAPE.OVAL)
    tfn = nc.text_frame; tfn.vertical_anchor = MSO_ANCHOR.MIDDLE
    pn = tfn.paragraphs[0]; pn.alignment = PP_ALIGN.CENTER
    rn = pn.add_run(); rn.text = ar_num(i + 1); style_run(rn, 17, WHITE, True)
    textbox(s, x + 0.25, y + 0.1, col_w - 1.2, 0.4,
            [P([(t, 13, ac, True)], align=PP_ALIGN.RIGHT)])
    textbox(s, x + 0.25, y + 0.5, col_w - 1.2, 0.52,
            [P([(d, 10, INK)], align=PP_ALIGN.RIGHT, line=1.12)])
# final hint card in the empty 8th cell
x, y = MARGIN, 1.85 + 3 * 1.24
hint = rect(s, x, y, col_w, 1.08, fill=tint(GOLD, 0.88), line=GOLD, line_w=1.2,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16)
textbox(s, x + 0.25, y + 0.12, col_w - 0.5, 0.85,
        [P([("⏱ الزمن المتوقع: ", 11.5, GOLD, True),
            ("أسبوعان إلى أربعة أسابيع من العمل المنظّم تكفي غالباً للوصول إلى "
             "فجوة موثقة وقابلة للدفاع عنها أمام أي لجنة.", 10.5, INK)],
           align=PP_ALIGN.RIGHT, line=1.18)])
add_notes(s, "خارطة طريق تنفيذية من سبع خطوات (إضافة توسعية على محتوى "
             "الورشة): (1) تضييق المجال — كلما ضاق التخصص سهُل حصر أدبياته. "
             "(2) جمع 20–30 دراسة حديثة (آخر 5 سنوات) لضمان مواكبة الحالة "
             "الراهنة للمعرفة. (3) قراءة الخاتمات وحدود الدراسات أولاً — "
             "اختصار هائل للوقت. (4) بناء مصفوفة أدبيات: صفوف = الدراسات، "
             "أعمدة = العينة/السياق/المنهج/الزمن/النتائج. (5) تظليل الفراغات "
             "والتناقضات في المصفوفة — الفجوة تظهر بصرياً. (6) التحقق بالأدوات "
             "الذكية أن الفجوة ما زالت قائمة ولم تُسد بدراسة حديثة. (7) الصياغة "
             "النهائية بجملة محددة. المدة الواقعية للعملية كاملة: 2–4 أسابيع "
             "من العمل المنتظم.")

# ================================================================ 12 — formulation
s = new_slide()
header(s, "✍️", "صياغة الفجوة البحثية باحترافية", "Writing Your Gap Statement", BLUE)
# template card
rect(s, MARGIN, 1.85, CW, 2.05, fill=NAVY, line=GOLD, line_w=1.3,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
textbox(s, MARGIN + 0.4, 2.02, CW - 0.8, 0.4,
        [P([("القالب الجاهز — انسخه واملأ الفراغات:", 12.5, GOLD, True)],
           align=PP_ALIGN.RIGHT)])
textbox(s, MARGIN + 0.4, 2.5, CW - 0.8, 1.3,
        [P([("«على الرغم من أن الدراسات السابقة تناولت ", 15, WHITE),
            ("[الموضوع X]", 15, GOLD, True),
            ("، إلا أن ", 15, WHITE),
            ("[الجانب Y]", 15, TEAL, True),
            (" لم يُدرس بعد في ", 15, WHITE),
            ("[السياق Z]", 15, CORAL, True),
            ("، مما يستدعي إجراء ", 15, WHITE),
            ("[نوع الدراسة المقترحة]", 15, PURPLE, True),
            (".»", 15, WHITE)],
           align=PP_ALIGN.RIGHT, line=1.45)])
# worked example
rect(s, MARGIN, 4.1, CW, 1.0, fill=tint(TEAL, 0.92), line=TEAL, line_w=1.1,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
textbox(s, MARGIN + 0.35, 4.2, CW - 0.7, 0.85,
        [P([("مثال مُطبَّق: ", 12, TEAL, True),
            ("«على الرغم من أن الدراسات السابقة تناولت التعلم الإلكتروني في الجامعات، إلا أن أثره على "
             "القلق الأكاديمي لم يُدرس بعد لدى طلاب الجامعات العربية بعد 2020، مما يستدعي إجراء دراسة "
             "مختلطة (كمية ونوعية) لسد هذه الفجوة.»", 11.5, INK)],
           align=PP_ALIGN.RIGHT, line=1.3)])
# checklist chips
textbox(s, MARGIN, 5.35, CW, 0.4,
        [P([("قائمة الفحص النهائية قبل اعتماد الفجوة:", 13, INK, True)],
           align=PP_ALIGN.RIGHT)])
checks = [("موثّقة؟", "تستند لمراجعة فعلية", PINK),
          ("محددة؟", "موضوع وفئة وسياق وزمن", GOLD),
          ("مهمة؟", "سدّها يُحدث فرقاً", TEAL),
          ("قابلة للمعالجة؟", "منهجية واقعية متاحة", PURPLE),
          ("جديدة؟", "لم تُسدّ بدراسة أحدث", CORAL)]
chw = (CW - 4 * 0.22) / 5
for i, (c, d, ac) in enumerate(checks):
    x = 13.333 - MARGIN - chw - i * (chw + 0.22)
    y = 5.85
    rect(s, x, y, chw, 0.92, fill=tint(ac, 0.9), line=ac, line_w=1.1,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
    textbox(s, x + 0.1, y + 0.1, chw - 0.2, 0.38,
            [P([("✔ " + c, 12, ac, True)], align=PP_ALIGN.CENTER)])
    textbox(s, x + 0.1, y + 0.48, chw - 0.2, 0.36,
            [P([(d, 9, MUTED)], align=PP_ALIGN.CENTER)])
add_notes(s, "الصياغة هي المرحلة التي تتحول فيها الملاحظة إلى فجوة علمية "
             "رسمية. القالب المعياري يتكون من أربعة عناصر: (X) ما تناولته "
             "الدراسات فعلاً — يثبت اطلاعك، (Y) الجانب الغائب — جوهر الفجوة، "
             "(Z) السياق المحدد — يمنح الفجوة حدوداً واقعية، ثم نوع الدراسة "
             "المقترحة — يحولها إلى خطة عمل. المثال المطبق يوضح كيف تُملأ "
             "الفراغات بمحتوى حقيقي. قبل الاعتماد النهائي مرر فجوتك على "
             "قائمة الفحص الخماسية: موثقة، محددة، مهمة، قابلة للمعالجة، "
             "وجديدة (لم تسدها دراسة أحدث). إن اجتازت الخمسة فهي جاهزة "
             "لتتصدر مقدمة مقترحك البحثي.")

# ================================================================ 13 — summary
s = new_slide(dark=True)
circle(s, -1.4, 4.4, 4.2, TEAL, 14)
circle(s, 11.4, -1.8, 4.2, PURPLE, 14)
circle(s, 12.1, 5.9, 1.6, GOLD, 20)
textbox(s, 1.5, 0.75, 10.333, 0.8,
        [P([("الخلاصة: رسائل تأخذها معك", 32, GOLD, True)], align=PP_ALIGN.CENTER)])
textbox(s, 1.5, 1.55, 10.333, 0.4,
        [P([("Key Takeaways", 13, TEAL, False, True)], align=PP_ALIGN.CENTER, rtl=False)])
keys = [
    ("الفجوة البحثية هي جواز مرورك", "إلى الأصالة العلمية والقبول والنشر والتمويل — لا بحث ناجح بلا فجوة واضحة", TEAL),
    ("اعرف الأنواع الستة وابحث بمنهجية", "معرفة، ممارسة، أدلة، سكان، منهجية، زمن — ولكل نوع سؤاله المفتاحي", GOLD),
    ("استعن بالذكاء الاصطناعي وتحقّق دائماً", "الأدوات تسرّع الاكتشاف، لكن الحكم النهائي للمصادر الأصلية ولعقلك النقدي", PURPLE),
    ("الفجوة القوية: محددة وموثقة ومهمة", "وقابلة للمعالجة — استخدم القالب الجاهز وقائمة الفحص قبل الاعتماد", CORAL),
]
for i, (t, d, ac) in enumerate(keys):
    y = 2.2 + i * 1.02
    rect(s, 1.7, y, 9.93, 0.88, fill=NAVY_SOFT, line=ac, line_w=1.1,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
    nc = rect(s, 1.7 + 9.93 - 0.78, y + 0.15, 0.58, 0.58, fill=ac,
              shape=MSO_SHAPE.OVAL)
    tfn = nc.text_frame; tfn.vertical_anchor = MSO_ANCHOR.MIDDLE
    pn = tfn.paragraphs[0]; pn.alignment = PP_ALIGN.CENTER
    rn = pn.add_run(); rn.text = ar_num(i + 1); style_run(rn, 15, WHITE, True)
    textbox(s, 2.0, y + 0.06, 8.85, 0.4,
            [P([(t, 13.5, ac, True)], align=PP_ALIGN.RIGHT)])
    textbox(s, 2.0, y + 0.45, 8.85, 0.38,
            [P([(d, 10.5, WHITE)], align=PP_ALIGN.RIGHT)])
rect(s, 4.42, 6.5, 4.5, 0.045, fill=GOLD)
textbox(s, 1.5, 6.65, 10.333, 0.45,
        [P([("شكراً لكم — المحتوى العلمي: ", 12, WHITE),
            ("Dr. Hana Hossen Elahemer", 12, GOLD, True),
            ("  •  نسخة موسّعة ومعاد تصميمها", 12, WHITE)],
           align=PP_ALIGN.CENTER)])
add_notes(s, "الخلاصة في أربع رسائل: (1) الفجوة البحثية ليست ترفاً أكاديمياً "
             "بل شرط الأصالة وبوابة القبول والنشر والتمويل. (2) الأنواع الستة "
             "إطار عملي — استحضر السؤال المفتاحي لكل نوع أثناء القراءة. "
             "(3) أدوات الذكاء الاصطناعي مساعد قوي لكنها لا تعفيك من التحقق "
             "من المصادر الأصلية ومن إعمال تفكيرك النقدي. (4) قبل اعتماد أي "
             "فجوة: تأكد أنها محددة وموثقة ومهمة وقابلة للمعالجة، واستخدم "
             "قالب الصياغة وقائمة الفحص. بالتوفيق في رحلتك البحثية!")

# ---------------------------------------------------------------- save
core = prs.core_properties
core.title = "الفجوات البحثية — Research Gaps"
core.author = "Dr. Hana Hossen Elahemer (original content)"
core.subject = "دليل شامل لاكتشاف الفجوات البحثية — نسخة موسعة"
core.comments = "Rebuilt & expanded from the original workshop video"

OUT = "Research-Gaps-الفجوات-البحثية.pptx"
prs.save(OUT)
print(f"Saved: {OUT} — {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
